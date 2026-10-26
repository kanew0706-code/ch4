"""Excel 紀錄：每處理一張照片就在最後加一列。

防護 4（公式注入）在這個檔案的 append_record() 裡。
"""

import os
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

SHEET_TITLE = "紀錄"
HEADERS = ["日期時間", "上傳者", "地方名稱", "景物", "地方特色", "照片存放位置"]
DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"
CELL_MAX_CHARS = 500

# Excel 會把這些字開頭的內容當成公式（Tab、換行開頭在某些情況下也會）
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")
# openpyxl 不接受的控制字元；不清掉的話存檔會直接出錯
_ILLEGAL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


@dataclass(frozen=True)
class Record:
    taken_at: datetime
    uploader: str
    place_name: str
    scenery: str
    intro: str
    photo_path: str


def escape_formula(text: str) -> str:
    """以 = + - @ 開頭的文字前面加一個 '，讓 Excel 當成純文字。"""
    return "'" + text if text.startswith(FORMULA_PREFIXES) else text


def append_record(path: Path, record: Record) -> None:
    """在 Excel 最後加一列。

    檔案被 Excel 打開而寫不進去時會丟 PermissionError，由呼叫的人決定要不要重試。
    """
    workbook = _open_or_create(path)
    sheet = workbook.active

    texts = [
        _clean(record.uploader),
        _clean(record.place_name),
        _clean(record.scenery),
        _clean(record.intro),
        _clean(record.photo_path),
    ]
    # 沒有防護時：文字原樣寫入，= 開頭的內容會被 Excel 當成公式執行
    # ==== 防護 4：Excel 公式注入　開始 ====
    # 上傳者名稱、地方名稱、景物、特色都可能被別人控制。以 = + - @ 開頭的，前面加 ' 變成純文字。
    texts = [escape_formula(text) for text in texts]
    # ==== 防護 4 結束 ====

    sheet.append([record.taken_at, *texts])
    sheet.cell(row=sheet.max_row, column=1).number_format = DATETIME_FORMAT

    # 先存成暫存檔再換掉原檔：存到一半當掉也不會把原本的紀錄弄壞
    temp_path = path.with_name(path.name + ".tmp")
    workbook.save(temp_path)
    try:
        os.replace(temp_path, path)
    except PermissionError:
        temp_path.unlink(missing_ok=True)
        raise


def _open_or_create(path: Path) -> Workbook:
    if path.exists():
        return load_workbook(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_TITLE
    sheet.append(HEADERS)
    return workbook


def _clean(text: str) -> str:
    """去掉 Excel 不接受的控制字元、限制長度。這是一般的資料整理，不是防護。"""
    return _ILLEGAL_CHARS.sub("", text)[:CELL_MAX_CHARS]
