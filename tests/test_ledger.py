from datetime import datetime

from openpyxl import load_workbook

from place_agent.ledger import HEADERS, Record, append_record


def record(uploader: str = "小明", place_name: str = "台北 101") -> Record:
    return Record(datetime(2026, 10, 26, 12, 0), uploader, place_name, "摩天大樓", "很高。", "photos/x.jpg")


def test_appends_rows_without_overwriting(tmp_path):
    path = tmp_path / "records.xlsx"
    append_record(path, record(place_name="台北 101"))
    append_record(path, record(place_name="不是地方"))

    rows = list(load_workbook(path).active.iter_rows(values_only=True))
    assert list(rows[0]) == HEADERS
    assert [row[2] for row in rows[1:]] == ["台北 101", "不是地方"]


def test_formula_like_text_is_stored_as_plain_text(tmp_path):
    path = tmp_path / "records.xlsx"
    append_record(path, record(uploader="=1+1", place_name="@SUM(A1)"))

    cells = list(load_workbook(path).active.iter_rows(min_row=2))[0]
    assert cells[1].value == "'=1+1"
    assert cells[1].data_type == "s"  # 字串，不是公式
    assert cells[2].value == "'@SUM(A1)"


def test_control_characters_do_not_break_saving(tmp_path):
    path = tmp_path / "records.xlsx"
    append_record(path, record(uploader="壞\x00字\x07元"))
    assert load_workbook(path).active.cell(row=2, column=2).value == "壞字元"
